"""
Report Downloads - Funciones de vista simples para descarga de reportes
Enfoque directo usando @api_view en lugar de clases APIView
"""
import csv
import io
from datetime import datetime
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated, IsAdminUser
from rest_framework.response import Response
from rest_framework import status


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsAdminUser])
def download_bookings_report(request):
    """
    Descarga reporte de reservas en CSV, PDF o Excel
    GET /api/downloads/bookings/?format=csv
    """
    from bookings.models import Booking
    
    format_type = request.GET.get('format', 'csv').lower()
    
    # Obtener datos
    bookings = Booking.objects.all().select_related('parent', 'nanny').order_by('-created_at')[:500]
    
    if format_type == 'csv':
        return _generate_bookings_csv(bookings)
    elif format_type == 'pdf':
        return _generate_bookings_pdf(bookings)
    elif format_type == 'excel':
        return _generate_bookings_excel(bookings)
    else:
        return Response({"error": "Formato no válido. Use: csv, pdf, excel"}, status=400)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsAdminUser])
def download_users_report(request):
    """
    Descarga reporte de usuarios en CSV, PDF o Excel
    GET /api/downloads/users/?format=csv
    """
    from django.contrib.auth import get_user_model
    User = get_user_model()
    
    format_type = request.GET.get('format', 'csv').lower()
    users = User.objects.all().order_by('-date_joined')[:500]
    
    if format_type == 'csv':
        return _generate_users_csv(users)
    elif format_type == 'pdf':
        return _generate_users_pdf(users)
    elif format_type == 'excel':
        return _generate_users_excel(users)
    else:
        return Response({"error": "Formato no válido. Use: csv, pdf, excel"}, status=400)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsAdminUser])
def download_payments_report(request):
    """
    Descarga reporte de pagos en CSV, PDF o Excel
    GET /api/downloads/payments/?format=csv
    """
    from bookings.models import Booking
    
    format_type = request.GET.get('format', 'csv').lower()
    bookings = Booking.objects.filter(status='completed').select_related('parent', 'nanny').order_by('-updated_at')[:500]
    
    if format_type == 'csv':
        return _generate_payments_csv(bookings)
    elif format_type == 'pdf':
        return _generate_payments_pdf(bookings)
    elif format_type == 'excel':
        return _generate_payments_excel(bookings)
    else:
        return Response({"error": "Formato no válido. Use: csv, pdf, excel"}, status=400)


# ============ GENERADORES CSV ============

def _generate_bookings_csv(bookings):
    """Genera CSV de reservas"""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="reservas_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'
    response.write('\ufeff')  # BOM para Excel
    
    writer = csv.writer(response)
    writer.writerow(['ID', 'Padre', 'Niñera', 'Estado', 'Fecha Inicio', 'Fecha Fin', 'Creado'])
    
    for b in bookings:
        writer.writerow([
            b.id,
            b.parent.username if b.parent else 'N/A',
            b.nanny.username if b.nanny else 'N/A',
            b.status,
            b.start_time.strftime('%Y-%m-%d %H:%M') if b.start_time else '',
            b.end_time.strftime('%Y-%m-%d %H:%M') if b.end_time else '',
            b.created_at.strftime('%Y-%m-%d %H:%M') if b.created_at else ''
        ])
    
    return response


def _generate_users_csv(users):
    """Genera CSV de usuarios"""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="usuarios_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'
    response.write('\ufeff')
    
    writer = csv.writer(response)
    writer.writerow(['ID', 'Usuario', 'Email', 'Rol', 'Activo', 'Staff', 'Fecha Registro'])
    
    for u in users:
        # Determinar rol
        role = 'admin' if u.is_staff else ('nanny' if hasattr(u, 'nanny_profile') else 'parent')
        writer.writerow([
            u.id,
            u.username,
            u.email,
            role,
            'Sí' if u.is_active else 'No',
            'Sí' if u.is_staff else 'No',
            u.date_joined.strftime('%Y-%m-%d %H:%M') if u.date_joined else ''
        ])
    
    return response


def _generate_payments_csv(bookings):
    """Genera CSV de pagos (reservas completadas)"""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="pagos_{datetime.now().strftime("%Y%m%d_%H%M")}.csv"'
    response.write('\ufeff')
    
    writer = csv.writer(response)
    writer.writerow(['ID Reserva', 'Padre', 'Niñera', 'Fecha Servicio', 'Estado', 'Completado'])
    
    for b in bookings:
        writer.writerow([
            b.id,
            b.parent.username if b.parent else 'N/A',
            b.nanny.username if b.nanny else 'N/A',
            b.start_time.strftime('%Y-%m-%d') if b.start_time else '',
            b.status,
            b.updated_at.strftime('%Y-%m-%d %H:%M') if b.updated_at else ''
        ])
    
    return response


# ============ GENERADORES PDF ============

def _generate_bookings_pdf(bookings):
    """Genera PDF de reservas"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
        elements = []
        
        styles = getSampleStyleSheet()
        elements.append(Paragraph(f"Reporte de Reservas - {datetime.now().strftime('%Y-%m-%d')}", styles['Heading1']))
        
        # Datos
        data = [['ID', 'Padre', 'Niñera', 'Estado', 'Fecha Inicio', 'Fecha Fin']]
        for b in bookings[:100]:  # Limitar para PDF
            data.append([
                str(b.id),
                b.parent.username if b.parent else 'N/A',
                b.nanny.username if b.nanny else 'N/A',
                b.status,
                b.start_time.strftime('%Y-%m-%d') if b.start_time else '',
                b.end_time.strftime('%Y-%m-%d') if b.end_time else ''
            ])
        
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(table)
        
        doc.build(elements)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="reservas_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
        
    except ImportError:
        return Response({"error": "reportlab no instalado"}, status=500)


def _generate_users_pdf(users):
    """Genera PDF de usuarios"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        
        styles = getSampleStyleSheet()
        elements.append(Paragraph(f"Reporte de Usuarios - {datetime.now().strftime('%Y-%m-%d')}", styles['Heading1']))
        
        data = [['ID', 'Usuario', 'Email', 'Rol', 'Activo']]
        for u in users[:100]:
            role = 'admin' if u.is_staff else ('nanny' if hasattr(u, 'nanny_profile') else 'parent')
            data.append([str(u.id), u.username, u.email or '', role, 'Sí' if u.is_active else 'No'])
        
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkblue),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(table)
        
        doc.build(elements)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="usuarios_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
        
    except ImportError:
        return Response({"error": "reportlab no instalado"}, status=500)


def _generate_payments_pdf(bookings):
    """Genera PDF de pagos"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
        from reportlab.lib.styles import getSampleStyleSheet
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        elements = []
        
        styles = getSampleStyleSheet()
        elements.append(Paragraph(f"Reporte de Pagos - {datetime.now().strftime('%Y-%m-%d')}", styles['Heading1']))
        
        data = [['ID', 'Padre', 'Niñera', 'Fecha', 'Estado']]
        for b in bookings[:100]:
            data.append([
                str(b.id),
                b.parent.username if b.parent else 'N/A',
                b.nanny.username if b.nanny else 'N/A',
                b.start_time.strftime('%Y-%m-%d') if b.start_time else '',
                b.status
            ])
        
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.darkgreen),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(table)
        
        doc.build(elements)
        buffer.seek(0)
        
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="pagos_{datetime.now().strftime("%Y%m%d")}.pdf"'
        return response
        
    except ImportError:
        return Response({"error": "reportlab no instalado"}, status=500)


# ============ GENERADORES EXCEL ============

def _generate_bookings_excel(bookings):
    """Genera Excel de reservas"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Reservas"
        
        # Headers
        headers = ['ID', 'Padre', 'Niñera', 'Estado', 'Fecha Inicio', 'Fecha Fin', 'Creado']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        
        # Datos
        for row, b in enumerate(bookings, 2):
            ws.cell(row=row, column=1, value=b.id)
            ws.cell(row=row, column=2, value=b.parent.username if b.parent else 'N/A')
            ws.cell(row=row, column=3, value=b.nanny.username if b.nanny else 'N/A')
            ws.cell(row=row, column=4, value=b.status)
            ws.cell(row=row, column=5, value=b.start_time.strftime('%Y-%m-%d %H:%M') if b.start_time else '')
            ws.cell(row=row, column=6, value=b.end_time.strftime('%Y-%m-%d %H:%M') if b.end_time else '')
            ws.cell(row=row, column=7, value=b.created_at.strftime('%Y-%m-%d %H:%M') if b.created_at else '')
        
        # Ajustar anchos
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 15
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="reservas_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
        
    except ImportError:
        return Response({"error": "openpyxl no instalado"}, status=500)


def _generate_users_excel(users):
    """Genera Excel de usuarios"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Usuarios"
        
        headers = ['ID', 'Usuario', 'Email', 'Rol', 'Activo', 'Staff', 'Fecha Registro']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
        
        for row, u in enumerate(users, 2):
            role = 'admin' if u.is_staff else ('nanny' if hasattr(u, 'nanny_profile') else 'parent')
            ws.cell(row=row, column=1, value=u.id)
            ws.cell(row=row, column=2, value=u.username)
            ws.cell(row=row, column=3, value=u.email or '')
            ws.cell(row=row, column=4, value=role)
            ws.cell(row=row, column=5, value='Sí' if u.is_active else 'No')
            ws.cell(row=row, column=6, value='Sí' if u.is_staff else 'No')
            ws.cell(row=row, column=7, value=u.date_joined.strftime('%Y-%m-%d') if u.date_joined else '')
        
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 15
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="usuarios_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
        
    except ImportError:
        return Response({"error": "openpyxl no instalado"}, status=500)


def _generate_payments_excel(bookings):
    """Genera Excel de pagos"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Pagos"
        
        headers = ['ID Reserva', 'Padre', 'Niñera', 'Fecha Servicio', 'Estado', 'Completado']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="FF6F00", end_color="FF6F00", fill_type="solid")
        
        for row, b in enumerate(bookings, 2):
            ws.cell(row=row, column=1, value=b.id)
            ws.cell(row=row, column=2, value=b.parent.username if b.parent else 'N/A')
            ws.cell(row=row, column=3, value=b.nanny.username if b.nanny else 'N/A')
            ws.cell(row=row, column=4, value=b.start_time.strftime('%Y-%m-%d') if b.start_time else '')
            ws.cell(row=row, column=5, value=b.status)
            ws.cell(row=row, column=6, value=b.updated_at.strftime('%Y-%m-%d') if b.updated_at else '')
        
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 15
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="pagos_{datetime.now().strftime("%Y%m%d")}.xlsx"'
        return response
        
    except ImportError:
        return Response({"error": "openpyxl no instalado"}, status=500)
